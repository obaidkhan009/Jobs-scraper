"""
Excel exporter for job listings with formatting and hyperlinks.
"""

import os
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scrapers.base_scraper import Job


class ExcelExporter:
    """Exports job listings to a formatted Excel file."""
    
    # Color scheme
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
    
    HIGH_SCORE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    MED_SCORE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    LOW_SCORE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    LINK_FONT = Font(color="0563C1", underline="single")
    
    # Column configuration
    COLUMNS = [
        ("Job Title", 40),
        ("Company", 25),
        ("Location", 20),
        ("Salary", 20),
        ("Good Tech Found", 50),
        ("Bad Tech Found", 30),
        ("Bad Tech Count", 12),
        ("Score", 10),
        ("Platform", 15),
        ("Date Posted", 15),
        ("Job Link", 50),
    ]
    
    def __init__(self, output_dir: str = "output"):
        """
        Initialize the Excel exporter.
        
        Args:
            output_dir: Directory to save Excel files
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def export(self, jobs: List[Job], filename: str = None) -> str:
        """
        Export jobs to an Excel file.
        
        Args:
            jobs: List of Job objects to export
            filename: Optional custom filename (without extension)
            
        Returns:
            Path to the created Excel file
        """
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"jobs_{timestamp}"
        
        filepath = os.path.join(self.output_dir, f"{filename}.xlsx")
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Job Listings"
        
        # Set up headers
        self._write_headers(ws)
        
        # Write job data
        for row_idx, job in enumerate(jobs, start=2):
            self._write_job_row(ws, row_idx, job)
        
        # Apply formatting
        self._apply_formatting(ws, len(jobs))
        
        # Add summary sheet
        self._add_summary_sheet(wb, jobs)
        
        # Save the workbook
        wb.save(filepath)
        
        print(f"[ExcelExporter] Exported {len(jobs)} jobs to {filepath}")
        return filepath
    
    def _write_headers(self, ws):
        """Write column headers with formatting."""
        for col_idx, (header, width) in enumerate(self.COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        
        # Freeze the header row
        ws.freeze_panes = "A2"
    
    def _write_job_row(self, ws, row_idx: int, job: Job):
        """Write a single job row to the worksheet."""
        # Job Title
        ws.cell(row=row_idx, column=1, value=job.title)
        
        # Company
        ws.cell(row=row_idx, column=2, value=job.company)
        
        # Location
        ws.cell(row=row_idx, column=3, value=job.location)
        
        # Salary
        ws.cell(row=row_idx, column=4, value=job.salary or "Not specified")
        
        # Good Tech Found
        good_tech_str = ", ".join(job.good_tech_found) if job.good_tech_found else "None"
        ws.cell(row=row_idx, column=5, value=good_tech_str)
        
        # Bad Tech Found
        bad_tech_str = ", ".join(job.bad_tech_found) if job.bad_tech_found else "None"
        ws.cell(row=row_idx, column=6, value=bad_tech_str)
        
        # Bad Tech Count
        ws.cell(row=row_idx, column=7, value=len(job.bad_tech_found))
        
        # Score
        score_cell = ws.cell(row=row_idx, column=8, value=job.score)
        
        # Apply color based on score
        if job.score >= 80:
            score_cell.fill = self.HIGH_SCORE_FILL
        elif job.score >= 50:
            score_cell.fill = self.MED_SCORE_FILL
        else:
            score_cell.fill = self.LOW_SCORE_FILL
        
        # Platform
        ws.cell(row=row_idx, column=9, value=job.platform)
        
        # Date Posted
        ws.cell(row=row_idx, column=10, value=job.date_posted or "Unknown")
        
        # Job Link (as hyperlink)
        if job.url:
            link_cell = ws.cell(row=row_idx, column=11)
            link_cell.value = job.url
            link_cell.hyperlink = job.url
            link_cell.font = self.LINK_FONT
        else:
            ws.cell(row=row_idx, column=11, value="No link available")
    
    def _apply_formatting(self, ws, num_jobs: int):
        """Apply general formatting to the worksheet."""
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Apply borders and alignment to all data cells
        for row in ws.iter_rows(min_row=1, max_row=num_jobs + 1, 
                                min_col=1, max_col=len(self.COLUMNS)):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:  # Skip header row
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        
        # Set row height for data rows
        for row_idx in range(2, num_jobs + 2):
            ws.row_dimensions[row_idx].height = 25
    
    def _add_summary_sheet(self, wb, jobs: List[Job]):
        """Add a summary sheet with statistics."""
        ws = wb.create_sheet(title="Summary")
        
        # Calculate statistics
        total_jobs = len(jobs)
        high_score_jobs = len([j for j in jobs if j.score >= 80])
        med_score_jobs = len([j for j in jobs if 50 <= j.score < 80])
        low_score_jobs = len([j for j in jobs if j.score < 50])
        
        # Unique companies
        unique_companies = len(set(j.company for j in jobs))
        
        # Average score
        avg_score = sum(j.score for j in jobs) / total_jobs if total_jobs else 0
        
        # Write summary
        summary_data = [
            ("Job Search Summary", ""),
            ("", ""),
            ("Generated On", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("", ""),
            ("Total Jobs Found", total_jobs),
            ("Unique Companies", unique_companies),
            ("Average Score", f"{avg_score:.1f}"),
            ("", ""),
            ("Score Distribution", ""),
            ("High Score (80+)", high_score_jobs),
            ("Medium Score (50-79)", med_score_jobs),
            ("Low Score (<50)", low_score_jobs),
        ]
        
        for row_idx, (label, value) in enumerate(summary_data, start=1):
            ws.cell(row=row_idx, column=1, value=label)
            ws.cell(row=row_idx, column=2, value=value)
            
            if row_idx == 1:
                ws.cell(row=row_idx, column=1).font = Font(bold=True, size=14)
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
